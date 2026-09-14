"""Xuất kết quả chấm bài ra file Excel."""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter


def export_results_to_excel(rows, out_path):
    """rows: danh sách tuple từ ResultStore.all_results()
    (id, sbd, exam_code, part1_score, part2_score, part3_score, total, source_image, created_at)
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "KetQua"

    headers = ["STT", "Số báo danh", "Mã đề", "Điểm phần I", "Điểm phần II", "Điểm phần III",
               "TỔNG ĐIỂM", "File ảnh", "Thời gian chấm"]
    ws.append(headers)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for i, row in enumerate(rows, start=1):
        _id, sbd, exam_code, p1, p2, p3, total, source_image, created_at = row
        ws.append([i, sbd, exam_code, p1, p2, p3, total, source_image, created_at])

    widths = [6, 14, 10, 12, 12, 13, 12, 30, 20]
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w

    wb.save(out_path)
    return out_path
